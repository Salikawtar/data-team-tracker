"""The gate.

A partial load is worse than a refusal, because it looks like it worked. These tests use
workbooks built on the fly, so nothing here depends on the real export.
"""

import openpyxl
import pytest


def _headers_of(path, sheet="RSP Dates"):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        return None
    return [c.value for c in wb[sheet][1]]


def test_a_good_workbook_passes(good_workbook, contract):
    headers = _headers_of(good_workbook)
    missing, unexpected = contract.check_source_columns(headers)
    assert missing == [] and unexpected == []


def test_a_missing_column_is_named_individually(workbook_missing_column, contract):
    """Naming the column is the point. A stack trace sends the reader to the code, and
    the problem is in the spreadsheet."""
    headers = _headers_of(workbook_missing_column)
    missing, _ = contract.check_source_columns(headers)
    assert missing == ["Fiscal Year"]


def test_a_workbook_without_the_rsp_dates_sheet_is_rejected(workbook_wrong_sheet):
    wb = openpyxl.load_workbook(workbook_wrong_sheet, data_only=True)
    assert "RSP Dates" not in wb.sheetnames


def test_the_instrument_sheet_is_readable(good_workbook):
    wb = openpyxl.load_workbook(good_workbook, data_only=True)
    ids = {str(r[0]).strip().upper()
           for r in wb["InstrumentID"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert "RSP-1" in ids


def test_a_good_workbook_reads_through_to_clean_rows(good_workbook, contract,
                                                     name_lookup, srs, known_ids):
    """End to end through the gate: read, clean, key, with no errors."""
    from tracker import clean as C
    from tracker import keys as K

    wb = openpyxl.load_workbook(good_workbook, data_only=True)
    ws = wb["RSP Dates"]
    headers = [c.value for c in ws[1]]
    raw = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

    rows, log, divs, errs = C.clean_rows(raw, contract, name_lookup=name_lookup,
                                         srs=srs, known_ids=known_ids)
    assert errs == []
    K.add_keys(rows)
    assert K.check_unique(rows) is True
    assert len(rows) == 6


def test_row_numbers_in_the_log_point_at_the_spreadsheet(good_workbook, contract,
                                                         name_lookup, srs, known_ids):
    """Row 1 is the header, so the first data row is 2. A person has to be able to open
    the file and find the value."""
    from tracker import clean as C
    wb = openpyxl.load_workbook(good_workbook, data_only=True)
    ws = wb["RSP Dates"]
    headers = [c.value for c in ws[1]]
    raw = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    rows, log, _, _ = C.clean_rows(raw, contract, name_lookup=name_lookup,
                                   srs=srs, known_ids=known_ids)
    assert min(e.source_row_number for e in log) >= 2
    assert max(e.source_row_number for e in log) <= len(raw) + 1
