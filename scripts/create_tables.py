"""Create the eight tracker tables.

    python scripts/create_tables.py [--drop] [--dry-run]

This script does not contain any table definitions. It reads them out of
conf/data_dictionary.xlsx and generates the DDL from it. That is the point: the dictionary
is the contract, and if a field only exists in a script then the contract is already
broken.

HOW IT WORKS

Read the field rows out of the dictionary, map each declared type to a SQL type, and emit
one CREATE TABLE per table with its primary and foreign keys declared inside it. Tables are
created parents first, so a foreign key always has something to point at.

The keys are enforced, not decorative. A bug that would otherwise surface as a wrong number
on a screen surfaces as an error at the moment of writing instead, which is the cheaper
place to find it.

The column comments come across too, so the documentation written in the dictionary is
attached to the tables rather than living in a file nobody opens.
"""

import argparse
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
for p in (os.path.join(PROJECT, "src"), os.path.join(PROJECT, "conf"),
          os.path.join(PROJECT, "app")):
    if p not in sys.path:
        sys.path.insert(0, p)

import openpyxl  # noqa: E402

import config  # noqa: E402
from lib import db  # noqa: E402


# ---------------------------------------------------------------- 1. the contract

SHEET_MAP = {
    "04_official_timeline":   ("core", "official_timeline"),
    "05_timeline_change_log": ("core", "timeline_change_log"),
    "06_import_run":          ("core", "import_run"),
    "07_data_work_items":     ("core", "data_work_items"),
    "08_team_updates":        ("core", "team_updates"),
}

# Foreign keys, written out because the dictionary names the referenced column but not the
# referenced table. Order of creation follows from this: a table is created after anything
# it points at.
FK_TARGETS = {
    ("official_timeline", "source_import_id"):     ("import_run", "import_id"),
    ("timeline_change_log", "timeline_item_id"):   ("official_timeline", "timeline_item_id"),
    ("timeline_change_log", "import_id"):          ("import_run", "import_id"),
    ("data_work_items", "timeline_item_id"):       ("official_timeline", "timeline_item_id"),
    ("team_updates", "data_work_item_id"):         ("data_work_items", "data_work_item_id"),
}

# Parents before children.
CORE_ORDER = ["import_run", "official_timeline", "timeline_change_log",
              "data_work_items", "team_updates"]

# One foreign key is declared in the dictionary but deliberately not enforced here.
#
# WHY. official_timeline is both a child (of import_run, through source_import_id) and a
# parent (of timeline_change_log and data_work_items). DuckDB implements an UPDATE on a
# table that has its own outgoing foreign key as a delete followed by an insert. That
# delete then trips the incoming foreign keys, so no row in official_timeline could ever be
# updated, which is the one thing the weekly import exists to do.
#
# Of the three keys, this is the one to give up. source_import_id is provenance: it says
# which import last touched a row, and import_run is written in the same transaction by the
# same script, so nothing else can put a stale value there. The other two are load bearing.
# data_work_items.timeline_item_id in particular is the single crossing between the official
# record and the Data Team's own work, and the whole ownership rule rests on it, so that one
# stays enforced.
SKIP_FK = {("official_timeline", "source_import_id")}

TYPE_MAP = {
    "string": "VARCHAR",
    "integer": "INTEGER",
    "date": "DATE",
    "timestamp": "TIMESTAMP",
    "boolean": "BOOLEAN",
}


def sql_type(dict_type):
    """'string(16)' and 'string' both become VARCHAR. Length is a validation rule, not a
    storage decision, and it is enforced in the pipeline where the message is readable."""
    base = re.sub(r"\(.*\)", "", str(dict_type)).strip().lower()
    if base not in TYPE_MAP:
        raise ValueError(f"unknown data_type {dict_type!r} in the dictionary")
    return TYPE_MAP[base]


def esc(text):
    return str(text or "").replace("'", "''").replace("\n", " ").strip()[:900]


def slug(header):
    s = re.sub(r"[^0-9a-zA-Z]+", "_", str(header).strip().lower()).strip("_")
    return re.sub(r"_+", "_", s)


def read_fields(wb, sheet_name):
    """One dict per field row, keyed by the dictionary's own column headers."""
    ws = wb[sheet_name]
    header = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        out.append({h: v for h, v in zip(header, row)})
    return out


# ---------------------------------------------------------------- 2. build the DDL


def build(wb):
    """Return (ddl_statements, comment_statements, table_names) in creation order."""
    fields = {s: read_fields(wb, s) for s in SHEET_MAP}

    missing = [s for s in SHEET_MAP if s not in wb.sheetnames]
    if missing:
        raise SystemExit(f"dictionary is missing sheets: {missing}")

    # Fail loudly on an unknown type rather than silently making it a VARCHAR.
    for rows in fields.values():
        for r in rows:
            sql_type(r["data_type"])

    ddl, comments, names = [], [], []

    # -- raw: one VARCHAR column per source column, because its whole job is to record
    #    what arrived before anyone touched it.
    ws = wb["12_Source_Profile"]
    source_headers = [r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1]]
    if len(source_headers) != 24:
        raise SystemExit(f"expected 24 source columns in sheet 12, found {len(source_headers)}")

    raw = config.t("raw", "timeline_import")
    raw_cols = ["  import_id VARCHAR NOT NULL", "  source_row_number INTEGER NOT NULL"]
    raw_cols += [f"  {slug(h)} VARCHAR" for h in source_headers]
    ddl.append((raw, f"CREATE TABLE IF NOT EXISTS {raw} (\n" + ",\n".join(raw_cols) + "\n)"))
    names.append(raw)
    comments.append(f"COMMENT ON TABLE {raw} IS '"
                    "Untouched copy of the RSP Dates sheet, one row per source row per import. "
                    "Every column is VARCHAR on purpose. Nothing in this table is ever corrected.'")
    comments.append(f"COMMENT ON COLUMN {raw}.import_id IS 'The import that produced this row'")
    comments.append(f"COMMENT ON COLUMN {raw}.source_row_number IS "
                    "'Row number in the RSP Dates sheet, 1 is the header'")
    for h in source_headers:
        comments.append(f"COMMENT ON COLUMN {raw}.{slug(h)} IS "
                        f"'Exactly as received from column \"{esc(h)}\"'")

    # -- clean: the official_timeline shape plus the import trace.
    clean = config.t("clean", "timeline_import")
    clean_cols = ["  import_id VARCHAR NOT NULL", "  source_row_number INTEGER NOT NULL"]
    for r in fields["04_official_timeline"]:
        if r["field_name"] in ("first_seen_date", "last_seen_date", "is_missing",
                               "source_import_id"):
            continue  # only meaningful once published
        clean_cols.append(f"  {r['field_name']} {sql_type(r['data_type'])}")
    ddl.append((clean, f"CREATE TABLE IF NOT EXISTS {clean} (\n" + ",\n".join(clean_cols) + "\n)"))
    names.append(clean)
    comments.append(f"COMMENT ON TABLE {clean} IS "
                    "'Typed, cleaned and keyed rows for one import. Input to the comparison.'")

    log = config.t("clean", "cleaning_log")
    ddl.append((log, f"""CREATE TABLE IF NOT EXISTS {log} (
  import_id VARCHAR NOT NULL,
  source_row_number INTEGER NOT NULL,
  natural_key VARCHAR,
  field_name VARCHAR NOT NULL,
  rule_id VARCHAR NOT NULL,
  raw_value VARCHAR,
  cleaned_value VARCHAR,
  cleaned_at TIMESTAMP NOT NULL
)"""))
    names.append(log)
    comments.append(f"COMMENT ON TABLE {log} IS 'One row per value the import corrected. This is "
                    "what keeps cleaning honest: the review shows a count and the reviewer can "
                    "open the detail.'")
    comments.append(f"COMMENT ON COLUMN {log}.natural_key IS "
                    "'The three key fields joined, for reading not for joining'")
    comments.append(f"COMMENT ON COLUMN {log}.rule_id IS "
                    "'The rule from sheet 10 of the dictionary, for example C05b'")

    # -- core: generated entirely from the dictionary.
    by_name = {name: sheet for sheet, (_, name) in SHEET_MAP.items()}
    for name in CORE_ORDER:
        sheet = by_name[name]
        rows = fields[sheet]
        full = config.t("core", name)

        cols, pk, table_comments = [], [], []
        for r in rows:
            col = r["field_name"]
            typ = sql_type(r["data_type"])
            nn = " NOT NULL" if str(r["nullable"]).strip().upper() == "N" else ""
            cols.append(f"  {col} {typ}{nn}")

            comment = esc(r.get("notes") or r.get("validation_rule"))
            if comment:
                table_comments.append(
                    f"COMMENT ON COLUMN {full}.{col} IS '{comment}'")

            key = str(r.get("key") or "").strip().upper()
            if key == "PK":
                pk.append(col)
            elif key == "FK":
                target = FK_TARGETS.get((name, col))
                if not target:
                    print(f"  note: {name}.{col} is marked FK but has no target mapped, skipping")
                    continue
                t_name, t_col = target
                if (name, col) in SKIP_FK:
                    ref = f"{config.t('core', t_name)}.{t_col}"
                    table_comments.append(
                        f"COMMENT ON COLUMN {full}.{col} IS "
                        f"'References {ref}. Declared in the data dictionary, deliberately "
                        f"not enforced here: see SKIP_FK in scripts/create_tables.py.'")
                    continue
                cols.append(f"  FOREIGN KEY ({col}) REFERENCES {config.t('core', t_name)} ({t_col})")

        if pk:
            # Declared after the columns and before the foreign keys would be tidier, but
            # DuckDB does not care about the order of table constraints.
            cols.append(f"  PRIMARY KEY ({', '.join(pk)})")

        ddl.append((full, f"CREATE TABLE IF NOT EXISTS {full} (\n" + ",\n".join(cols) + "\n)"))
        names.append(full)
        comments.append(
            f"COMMENT ON TABLE {full} IS '"
            f"Data Team Work Tracker. Defined by {sheet} of the data dictionary. "
            f"Do not add a column here without adding it there first.'")
        comments.extend(table_comments)

    return ddl, comments, names


# ---------------------------------------------------------------- 3. run it


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--dry-run", action="store_true",
                    help="print the DDL and create nothing")
    ap.add_argument("--drop", action="store_true",
                    help="drop the tables first. Destroys data.")
    args = ap.parse_args()

    print(config.describe())
    print("dictionary:", config.DICTIONARY)
    if not os.path.exists(config.DICTIONARY):
        raise SystemExit(f"Data dictionary not found at {config.DICTIONARY}")

    wb = openpyxl.load_workbook(config.DICTIONARY, data_only=True)
    ddl, comments, names = build(wb)

    if args.dry_run:
        for _, statement in ddl:
            print("-" * 78)
            print(statement + ";")
        print("-" * 78)
        print(f"\n{len(ddl)} tables, {len(comments)} comments. Nothing was created.")
        return

    if args.drop:
        print("\n--drop given. Dropping in reverse dependency order.")
        for full, _ in reversed(ddl):
            db.execute(f"DROP TABLE IF EXISTS {full}")
            print("  dropped", full)

    print()
    for full, statement in ddl:
        db.execute(statement)
        print("  created", full)

    for statement in comments:
        db.execute(statement)
    print(f"\n  {len(comments)} column and table comments applied")

    print("\nVerifying:\n")
    existing = {r[0] for _, rows in [db.execute("SHOW TABLES")] for r in rows}
    for full in names:
        mark = "ok  " if full in existing else "MISSING"
        print(f"  {mark} {full:<48} {db.table_count(full):>6} rows")

    gap = set(names) - existing
    if gap:
        raise SystemExit(f"\nMISSING: {sorted(gap)}")
    print(f"\nAll {len(names)} tables present. Run scripts/import_timeline.py next.")


if __name__ == "__main__":
    main()
